import os
from openpyxl import Workbook, load_workbook
from bambu_inventory.local_sheet import (
    LocalSheetWriter, HEADER, SHEET_NAME, INVENTORY_SHEET_NAME, INVENTORY_HEADER,
)


def test_ensure_workbook_creates_file_with_header_when_missing(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    writer = LocalSheetWriter(path)

    result = writer.ensure_workbook()

    assert result == path
    assert os.path.exists(path)
    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    assert [cell.value for cell in sheet[1]] == HEADER


def test_ensure_workbook_does_not_overwrite_existing_file(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    writer = LocalSheetWriter(path)
    writer.ensure_workbook()
    writer.append_rows([{
        "date": "2026-07-30T09:00:00Z", "print_name": "Vase", "printer": "X1C-A",
        "status": "finished", "duration_min": 30, "ams_slot": "0-1",
        "filament_type": "PLA", "colour_hex": "FF0000FF", "weight_g": 12.5,
        "task_id": "1002",
    }])

    writer2 = LocalSheetWriter(path)
    writer2.ensure_workbook()

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    assert sheet.max_row == 2


def test_append_rows_writes_values_in_header_order(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    writer = LocalSheetWriter(path)
    writer.ensure_workbook()
    row = {
        "date": "2026-07-30T09:00:00Z", "print_name": "Two Color Vase", "printer": "X1C-A",
        "status": "failed", "duration_min": 30, "ams_slot": "0-1",
        "filament_type": "PLA", "colour_hex": "FF0000FF", "weight_g": 12.5,
        "task_id": "1002",
    }

    writer.append_rows([row])

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    written = [cell.value for cell in sheet[2]]
    assert written == [
        "2026-07-30T09:00:00Z", "Two Color Vase", "X1C-A", "failed", 30,
        "0-1", "PLA", "FF0000FF", 12.5, "1002",
    ]


def test_append_rows_appends_multiple_calls_without_overwriting(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    writer = LocalSheetWriter(path)
    writer.ensure_workbook()
    row1 = {
        "date": "d", "print_name": "p", "printer": "pr", "status": "finished",
        "duration_min": 1, "ams_slot": "0-0", "filament_type": "PLA",
        "colour_hex": "000000FF", "weight_g": 1.0, "task_id": "1",
    }
    row2 = {**row1, "task_id": "2"}

    writer.append_rows([row1])
    writer.append_rows([row2])

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    assert sheet.max_row == 3


def test_append_rows_skips_row_with_already_logged_task_id(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    writer = LocalSheetWriter(path)
    writer.ensure_workbook()
    row = {
        "date": "d", "print_name": "p", "printer": "pr", "status": "finished",
        "duration_min": 1, "ams_slot": "0-0", "filament_type": "PLA",
        "colour_hex": "000000FF", "weight_g": 1.0, "task_id": "1",
    }

    writer.append_rows([row])
    writer.append_rows([row])

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    assert sheet.max_row == 2


def test_append_rows_skips_only_the_duplicate_when_mixed_with_new_rows(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    writer = LocalSheetWriter(path)
    writer.ensure_workbook()
    row1 = {
        "date": "d", "print_name": "p", "printer": "pr", "status": "finished",
        "duration_min": 1, "ams_slot": "0-0", "filament_type": "PLA",
        "colour_hex": "000000FF", "weight_g": 1.0, "task_id": "1",
    }
    row2 = {**row1, "task_id": "2"}

    writer.append_rows([row1])
    writer.append_rows([row1, row2])

    workbook = load_workbook(path)
    sheet = workbook[SHEET_NAME]
    assert sheet.max_row == 3
    task_ids = [cell.value for cell in sheet["J"][1:] if cell.value is not None]
    assert task_ids == ["1", "2"]


def _make_row(filament_type="PLA", colour_hex="FF0000FF", weight_g=10.0):
    return {
        "date": "d", "print_name": "p", "printer": "pr", "status": "finished",
        "duration_min": 1, "ams_slot": "0-0", "filament_type": filament_type,
        "colour_hex": colour_hex, "weight_g": weight_g, "task_id": "1",
    }


def test_sync_filament_inventory_writes_formula_into_matching_row(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet(INVENTORY_SHEET_NAME)
    sheet.append(INVENTORY_HEADER)
    sheet.append(["PLA", "FF0000FF", 1000, None])
    workbook.save(path)

    writer = LocalSheetWriter(path)
    writer.sync_filament_inventory([_make_row(weight_g=12.5)])

    result = load_workbook(path)
    sheet = result[INVENTORY_SHEET_NAME]
    formula = sheet.cell(row=2, column=4).value
    assert formula == "=C2-SUMIFS(PrintLog!I:I,PrintLog!G:G,A2,PrintLog!H:H,B2)"


def test_sync_filament_inventory_creates_row_when_no_match(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet(INVENTORY_SHEET_NAME)
    sheet.append(INVENTORY_HEADER)
    workbook.save(path)

    writer = LocalSheetWriter(path)
    writer.sync_filament_inventory([_make_row(filament_type="PETG", colour_hex="00FF00FF", weight_g=20.0)])

    result = load_workbook(path)
    sheet = result[INVENTORY_SHEET_NAME]
    assert sheet.max_row == 2
    new_row = [cell.value for cell in sheet[2]]
    assert new_row[0] == "PETG"
    assert new_row[1] == "00FF00FF"
    assert new_row[2] is None
    assert new_row[3] == "=C2-SUMIFS(PrintLog!I:I,PrintLog!G:G,A2,PrintLog!H:H,B2)"


def test_sync_filament_inventory_creates_inventory_sheet_if_missing(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    writer = LocalSheetWriter(path)
    writer.ensure_workbook()

    writer.sync_filament_inventory([_make_row(weight_g=5.0)])

    result = load_workbook(path)
    assert INVENTORY_SHEET_NAME in result.sheetnames
    sheet = result[INVENTORY_SHEET_NAME]
    assert [cell.value for cell in sheet[1]] == INVENTORY_HEADER
    row = [cell.value for cell in sheet[2]]
    assert row[0] == "PLA"
    assert row[1] == "FF0000FF"
    assert row[3] == "=C2-SUMIFS(PrintLog!I:I,PrintLog!G:G,A2,PrintLog!H:H,B2)"


def test_sync_filament_inventory_matches_only_exact_type_and_colour(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet(INVENTORY_SHEET_NAME)
    sheet.append(INVENTORY_HEADER)
    sheet.append(["PLA", "0000FFFF", 1000, None])
    workbook.save(path)

    writer = LocalSheetWriter(path)
    writer.sync_filament_inventory([_make_row(filament_type="PLA", colour_hex="FF0000FF", weight_g=10.0)])

    result = load_workbook(path)
    sheet = result[INVENTORY_SHEET_NAME]
    assert sheet.max_row == 3
    assert sheet.cell(row=2, column=4).value is None
    new_row = [cell.value for cell in sheet[3]]
    assert new_row[0] == "PLA"
    assert new_row[1] == "FF0000FF"
    assert new_row[3] == "=C3-SUMIFS(PrintLog!I:I,PrintLog!G:G,A3,PrintLog!H:H,B3)"


def test_sync_filament_inventory_refreshes_formula_on_existing_row(tmp_path):
    path = str(tmp_path / "inventory.xlsx")
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet(INVENTORY_SHEET_NAME)
    sheet.append(INVENTORY_HEADER)
    # simulates a row left over from the old numeric-deduction behaviour
    sheet.append(["PLA", "FF0000FF", 1000, -42])
    workbook.save(path)

    writer = LocalSheetWriter(path)
    writer.sync_filament_inventory([_make_row(weight_g=12.5)])

    result = load_workbook(path)
    sheet = result[INVENTORY_SHEET_NAME]
    assert sheet.cell(row=2, column=4).value == "=C2-SUMIFS(PrintLog!I:I,PrintLog!G:G,A2,PrintLog!H:H,B2)"
